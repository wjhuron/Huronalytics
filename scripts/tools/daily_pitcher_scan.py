#!/usr/bin/env python3
"""Daily pitcher scan: who stood out yesterday, and why.

Runs once per calendar day, after the first pipeline update of the day
(scripts/auto-pull.sh calls it when a pull moves data/metadata_rs.json and
the day's marker is absent). Reads the six division workbooks like the daily
card does, takes the latest game date in the data, and for every MLB or ROC
pitcher who threw that day reports two kinds of flag:

  1. CHANGE flags: a pitch type whose velocity, release point, extension,
     movement, spin, usage, whiff rate or chase rate sat outside his own
     game-to-game noise. z = delta / sqrt(day-to-day component + this
     outing's sampling noise); the day-to-day components were measured on
     2021-2025 as five independent replicates
     (scripts/research/cards/daily_delta_scale_v2.py, data/_daily_scan_scales.log).
     A NEW pitch = a type with zero prior pitches from him in the season
     data, on a day that is not his season debut.

  2. OUTING flags: the outing's whiff%, chase%, CSW%, Stuff+, Loc+, xRV/100
     and Pitching+ ranked against every 2026 outing at the same level, split
     at ROLE_SPLIT pitches into start-length and relief-length pools.

Output: ~/Downloads/daily_pitcher_scan_<date>.docx (the read) and
data/_daily_scan/scan_<date>.json (the data), plus the once-a-day marker
data/_daily_scan/last_run_date.

Usage:
  python3 scripts/tools/daily_pitcher_scan.py             # latest date in Sheets
  python3 scripts/tools/daily_pitcher_scan.py --date 2026-09-14
  python3 scripts/tools/daily_pitcher_scan.py --cached    # data/all_pitches_rs_cache.pkl
  python3 scripts/tools/daily_pitcher_scan.py --force     # ignore the day marker
"""
import os, sys, json, math, bisect, argparse, pickle
from datetime import date, datetime
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
from pipeline.utils import MLB_TEAMS, NON_PA_EVENTS  # noqa: E402

# ── Settings (per-run scratch; CLI flags override) ───────────────────────
SCAN_DATE = None      # 'YYYY-MM-DD' or None for the latest game date in the data
USE_CACHE = False     # True reads data/all_pitches_rs_cache.pkl instead of Sheets
FORCE = False         # True ignores the once-a-day marker
OUT_DIR = os.path.expanduser('~/Downloads')

# ── Constants ────────────────────────────────────────────────────────────
STATE_DIR = os.path.join(ROOT, 'data', '_daily_scan')
MARKER = os.path.join(STATE_DIR, 'last_run_date')
# Outing floor: 10 pitches, or an immaculate inning (9 pitches, 3 K).
MIN_OUTING = 10
# Baseline floor per pitch type before a delta is scored. The SE already
# prices a thin baseline (same argument as SEASON_DELTA_MIN in cards/pitcher.py).
BASELINE_MIN = 10
# Today's pitches of a type before a mean-column delta is scored (needs a
# within-outing SD).
TODAY_MIN = 5
# Flag threshold on the observed-delta z. A convention read off the
# persistence curves (data/_daily_scan_scales.log): at 2.0 about 4.6% of
# real starts flag per column, and a velocity flag persists into the next
# start 62% of the time against a 39% base; release X 83% vs 37%. The curve
# has no interior optimum: precision rises with t while recall falls, and
# there is no stated exchange rate between the two.
FLAG_Z = 2.0
# Pitcher-level rule (multiple comparisons: ~40 tests per pitcher at a 5%
# tail is two chance flags each; the single-type rule flagged 94 pitchers
# on 2026-09-13). A mean column flags when TWO pitch types moved the same
# way at |z| >= FLAG_Z, or one type moved at |z| >= SINGLE_Z. Measured on
# 2021-2025 pitcher-games (scratch rule_sweep, 2026-09-16): flags 3.6-5.3%
# of pitcher-games per column against 8-14% for the single-type rule, and
# the next start moves the same way 84% (release X), 58% (spin), 54% (velo),
# 53% (release Z), 42% (extension), 36-38% (movement) of the time, against
# 27-36% base rates. Usage uses the same counts with either sign (shares sum
# to one). Whiff and chase flag on SINGLE_Z only: they are hot-day flags and
# the outing side already ranks them. Convention, not an optimum.
SINGLE_Z = 3.0
# Extension flags persist at 41% against a 33% base at every t from 1.5 up
# (curve flat), so extension is reported with a LOW-CONFIDENCE label.
LOW_CONFIDENCE_COLS = {'ext'}
# Outing pool: percentile at or above this flags. Convention (see the
# flags-per-day table the run prints).
OUTING_PCTL = 95
# Pitcher-level outing rule: Pitching+ at or above OUTING_PCTL, or at least
# OUTING_MIN_HITS metrics at or above it, or K_HEADLINE strikeouts. Measured
# on 2026 MLB outings from June on (103 days, ~102 outings a day): any single
# metric at the 95th flags 24 outings a day, this rule 10.8, and the 98th
# with two hits 4.8. Convention.
OUTING_MIN_HITS = 2
# The document gives a full section to the top TOP_N pitchers by score and
# one line each to the rest. Convention: the change rule alone marks about
# 40 pitchers a day, which is a list, not a read.
TOP_N = 15
ROLE_SPLIT = 50        # pitches: >= is start-length, < is relief-length
POOL_MIN = 100         # outings per (level, role) pool before it is used
K_HEADLINE = 10        # strikeouts: an absolute headline regardless of pool
NEW_PITCH_MIN = 1

# Day-to-day scales, 2021-2025 pooled (daily_delta_scale_v2.py). Velo and
# usage are the card's shipped constants (cards/pitcher.py DAILY_*), which
# the v2 script reproduced (FF 0.669 vs 0.67).
SD_DAY = {
    'velo': {'FF': 0.67, 'SI': 0.65, 'FC': 0.81, 'SL': 0.92, 'ST': 0.80, 'CU': 0.80, 'CH': 0.77, 'FS': 0.80},
    'rx':   {'FF': 0.175, 'SI': 0.168, 'FC': 0.155, 'SL': 0.172, 'ST': 0.179, 'CU': 0.173, 'CH': 0.169, 'FS': 0.154},
    'rz':   {'FF': 0.090, 'SI': 0.094, 'FC': 0.087, 'SL': 0.090, 'ST': 0.095, 'CU': 0.088, 'CH': 0.090, 'FS': 0.097},
    'ext':  {'FF': 0.101, 'SI': 0.103, 'FC': 0.103, 'SL': 0.108, 'ST': 0.124, 'CU': 0.104, 'CH': 0.110, 'FS': 0.111},
    'ivb':  {'FF': 1.00, 'SI': 0.98, 'FC': 1.32, 'SL': 1.39, 'ST': 1.34, 'CU': 1.18, 'CH': 1.14, 'FS': 1.30},
    'hb':   {'FF': 1.07, 'SI': 1.15, 'FC': 1.02, 'SL': 1.09, 'ST': 1.51, 'CU': 1.28, 'CH': 1.28, 'FS': 1.52},
    'spin': {'FF': 49.6, 'SI': 51.3, 'FC': 52.9, 'SL': 55.6, 'ST': 60.4, 'CU': 64.5, 'CH': 72.4, 'FS': 97.7},
}
SD_DAY_DEFAULT = {'velo': 0.80, 'rx': 0.168, 'rz': 0.091, 'ext': 0.107,
                  'ivb': 1.21, 'hb': 1.24, 'spin': 60.4}
# Rate overdispersion c (day var = c * p(1-p)); CSW measured at c < 0 on
# every type and season, i.e. binomial only, so it takes 0.
RATE_C = {
    'usage': {'FF': 0.0215, 'SI': 0.0301, 'FC': 0.0279, 'SL': 0.0241, 'ST': 0.0266, 'CU': 0.0155, 'CH': 0.0160, 'FS': 0.0213},
    'whiff': {'FF': 0.0100, 'SI': 0.0041, 'FC': 0.0015, 'SL': 0.0067, 'ST': 0.0050, 'CU': 0.0060, 'CH': 0.0103, 'FS': 0.0179},
    'chase': {},   # per-type values straddle zero; pooled 0.0038 for every type
}
RATE_C_DEFAULT = {'usage': 0.0227, 'whiff': 0.0076, 'chase': 0.0038}

MEAN_COLS = [('velo', 'Velocity', 'mph', 1), ('rx', 'RelPosX', 'ft', 2), ('rz', 'RelPosZ', 'ft', 2),
             ('ext', 'Extension', 'ft', 2), ('ivb', 'IndVertBrk', 'in', 1), ('hb', 'HorzBrk', 'in', 1),
             ('spin', 'Spin Rate', 'rpm', 0)]
COL_LABEL = {'velo': 'velocity', 'rx': 'release X', 'rz': 'release height', 'ext': 'extension',
             'ivb': 'induced vertical break', 'hb': 'horizontal break', 'spin': 'spin rate',
             'usage': 'usage', 'whiff': 'whiff rate', 'chase': 'chase rate'}
PITCH_NAMES = {'FF': 'four-seamer', 'SI': 'sinker', 'FC': 'cutter', 'SL': 'slider', 'ST': 'sweeper',
               'CU': 'curveball', 'CH': 'changeup', 'FS': 'splitter', 'KC': 'knuckle-curve',
               'SV': 'slurve', 'CS': 'slow curve', 'KN': 'knuckleball', 'FO': 'forkball', 'SC': 'screwball'}
SWING_DESC = {'Swinging Strike', 'Foul', 'In Play'}   # bunts are not swings
WHIFF_DESC = {'Swinging Strike'}
CSW_DESC = {'Called Strike', 'Swinging Strike'}
OUTING_METRICS = [('whiff', 'Whiff%', True), ('chase', 'Chase%', True), ('csw', 'CSW%', True),
                  ('stuff', 'Stuff+', True), ('loc', 'Loc+', True), ('xrv100', 'xRV/100', True),
                  ('pplus', 'Pitching+', True)]


def sf(v):
    if v is None or v == '':
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(f) else f


def level_of(row):
    t = row.get('PTeam')
    if t in MLB_TEAMS:
        return 'MLB'
    if t == 'ROC':
        return 'ROC'
    return None


def load_pitches(use_cache):
    if use_cache:
        p = os.path.join(ROOT, 'data', 'all_pitches_rs_cache.pkl')
        print(f'reading {p}')
        with open(p, 'rb') as fh:
            rows = pickle.load(fh)
    else:
        from pipeline.fetch import read_all_pitches_from_sheets
        print('reading the six division workbooks (about 2m30s)')
        rows = read_all_pitches_from_sheets()
    rows = [r for r in rows if level_of(r) and r.get('Game Date') and r.get('Pitcher')]
    print(f'{len(rows)} MLB + ROC pitches, latest {max(r["Game Date"] for r in rows)}')
    return rows


# ── Per-outing metrics ───────────────────────────────────────────────────
def outing_counts(ps):
    n = len(ps)
    swings = [p for p in ps if p.get('Description') in SWING_DESC]
    ooz = [p for p in ps if p.get('InZone') == 'No']
    c = {
        'n': n,
        'bf': sum(1 for p in ps if p.get('Event') and p['Event'] not in NON_PA_EVENTS),
        'k': sum(1 for p in ps if str(p.get('Event') or '').startswith('Strikeout')),
        'bb': sum(1 for p in ps if p.get('Event') in ('Walk', 'Intent Walk')),
        'h': sum(1 for p in ps if p.get('Event') in ('Single', 'Double', 'Triple', 'Home Run')),
        'hr': sum(1 for p in ps if p.get('Event') == 'Home Run'),
        'swings': len(swings),
        'whiffs': sum(1 for p in swings if p.get('Description') in WHIFF_DESC),
        'ooz': len(ooz),
        'chases': sum(1 for p in ooz if p.get('Description') in SWING_DESC),
        'csws': sum(1 for p in ps if p.get('Description') in CSW_DESC),
    }
    c['whiff'] = c['whiffs'] / c['swings'] if c['swings'] else None
    c['chase'] = c['chases'] / c['ooz'] if c['ooz'] else None
    c['csw'] = c['csws'] / n if n else None
    for key, col in (('stuff', 'Stuff+'), ('loc', 'Loc+')):
        v = [sf(p.get(col)) for p in ps]
        v = [x for x in v if x is not None]
        c[key] = sum(v) / len(v) if v else None
    rv = [sf(p.get('RunExp')) for p in ps]
    rv = [x for x in rv if x is not None]
    c['rv'] = sum(rv) if rv else None
    xv = _xrv(ps)
    c['xrv100'] = sum(xv) / n * 100 if xv else None
    return c


def _xrv(ps):
    from cards.pitcher import _compute_pitch_xrv
    return _compute_pitch_xrv(ps)


def build_outing_pools(rows, scan_date):
    """(level, role) -> {metric: sorted values} over every 2026 outing before
    and including the scan date, plus the card's Pitching+ pool per level."""
    from cards.pitcher import _build_outing_pool, _outing_pitching_plus
    by_outing = defaultdict(list)
    for r in rows:
        if r['Game Date'] <= scan_date:
            by_outing[(level_of(r), r['Pitcher'], r['Game Date'])].append(r)
    pp_pool = {}
    for lvl in ('MLB', 'ROC'):
        lvl_rows = [r for r in rows if level_of(r) == lvl and r['Game Date'] <= scan_date]
        pp_pool[lvl] = _build_outing_pool(lvl_rows)
    pools = defaultdict(lambda: defaultdict(list))
    for (lvl, _p, _d), ps in by_outing.items():
        if len(ps) < MIN_OUTING:
            continue
        c = outing_counts(ps)
        role = 'SP' if c['n'] >= ROLE_SPLIT else 'RP'
        val, _pct = _outing_pitching_plus(ps, pp_pool[lvl])
        c['pplus'] = val
        for m, _lab, _hi in OUTING_METRICS:
            if c.get(m) is not None:
                pools[(lvl, role)][m].append(c[m])
    for k in pools:
        for m in pools[k]:
            pools[k][m].sort()
    return pools, pp_pool


def pctl(sorted_vals, x):
    if not sorted_vals:
        return None
    return int(round(bisect.bisect_left(sorted_vals, x) / len(sorted_vals) * 100))


# ── Change flags ─────────────────────────────────────────────────────────
def _mean_sd(vals):
    n = len(vals)
    if n == 0:
        return None, None
    m = sum(vals) / n
    if n < 2:
        return m, None
    var = sum((v - m) ** 2 for v in vals) / (n - 1)
    return m, math.sqrt(var)


def type_deltas(today, base, pt, hand_sign, tc_today, tc_base):
    """Per-type change record: value today, season baseline, z per column."""
    rec = {'pt': pt, 'n_today': len(today), 'n_base': len(base), 'cols': {}}
    for key, col, unit, dec in MEAN_COLS:
        # Sheet conventions (medians on the 2026 cache, 2026-09-16): RelPosX is
        # NEGATIVE for a RHP (arm side = negative), HorzBrk is POSITIVE for a
        # RHP sinker (arm side = positive). Both flip for a LHP. The scan
        # stores both as positive = arm side.
        sign = {'rx': -hand_sign, 'hb': hand_sign}.get(key, 1.0)
        t = [sf(p.get(col)) for p in today]; t = [x * sign for x in t if x is not None]
        b = [sf(p.get(col)) for p in base]; b = [x * sign for x in b if x is not None]
        if len(t) < TODAY_MIN or len(b) < BASELINE_MIN:
            continue
        mt, st = _mean_sd(t); mb, _ = _mean_sd(b)
        sd_day = SD_DAY[key].get(pt, SD_DAY_DEFAULT[key])
        se = math.sqrt(sd_day ** 2 + (st ** 2 / len(t) if st else 0.0))
        rec['cols'][key] = {'today': mt, 'base': mb,
                            'delta': mt - mb, 'z': (mt - mb) / se, 'unit': unit, 'dec': dec,
                            'signed': key in ('rx', 'hb')}
    # usage
    u_t, u_b = len(today) / tc_today, len(base) / tc_base
    if len(base) >= BASELINE_MIN and 0 < u_b < 1:
        c = RATE_C['usage'].get(pt, RATE_C_DEFAULT['usage'])
        se = math.sqrt(u_b * (1 - u_b) * (c + 1.0 / tc_today))
        rec['cols']['usage'] = {'today': u_t, 'base': u_b, 'delta': u_t - u_b, 'z': (u_t - u_b) / se,
                                'unit': '%', 'dec': 0, 'pct': True}
    # whiff, chase
    for key, num_desc, den_fn in (('whiff', WHIFF_DESC, lambda p: p.get('Description') in SWING_DESC),
                                  ('chase', SWING_DESC, lambda p: p.get('InZone') == 'No')):
        dt = [p for p in today if den_fn(p)]; db = [p for p in base if den_fn(p)]
        if len(dt) < TODAY_MIN or len(db) < BASELINE_MIN:
            continue
        p_t = sum(1 for p in dt if p.get('Description') in num_desc) / len(dt)
        p_b = sum(1 for p in db if p.get('Description') in num_desc) / len(db)
        if not (0.02 < p_b < 0.98):
            continue
        c = RATE_C[key].get(pt, RATE_C_DEFAULT[key])
        se = math.sqrt(p_b * (1 - p_b) * (c + 1.0 / len(dt)))
        rec['cols'][key] = {'today': p_t, 'base': p_b, 'delta': p_t - p_b, 'z': (p_t - p_b) / se,
                            'unit': '%', 'dec': 0, 'pct': True, 'n_den': len(dt)}
    return rec


# ── Scan ─────────────────────────────────────────────────────────────────
def scan(rows, scan_date):
    by_p = defaultdict(list)
    for r in rows:
        by_p[r['Pitcher']].append(r)
    pools, pp_pool = build_outing_pools(rows, scan_date)
    from cards.pitcher import _outing_pitching_plus
    pool_sizes = {f'{k[0]}_{k[1]}': len(v.get('csw', [])) for k, v in pools.items()}
    print('outing pools (outings):', pool_sizes)
    report = {'scan_date': scan_date, 'generated': datetime.now().isoformat(timespec='seconds'),
              'pool_sizes': pool_sizes, 'pitchers': [], 'skipped': []}
    for name, prs in by_p.items():
        today = [r for r in prs if r['Game Date'] == scan_date]
        if not today:
            continue
        lvl = level_of(today[0])
        base = [r for r in prs if r['Game Date'] < scan_date]
        cnt = outing_counts(today)
        immaculate = cnt['n'] == 9 and cnt['k'] == 3 and cnt['bf'] == 3
        if cnt['n'] < MIN_OUTING and not immaculate:
            report['skipped'].append({'pitcher': name, 'n': cnt['n']})
            continue
        team = today[0].get('PTeam'); opp = today[0].get('BTeam')
        hand = today[0].get('Throws') or (base[0].get('Throws') if base else None)
        hand_sign = -1.0 if hand == 'L' else 1.0
        role = 'SP' if cnt['n'] >= ROLE_SPLIT else 'RP'
        val, _ = _outing_pitching_plus(today, pp_pool[lvl])
        cnt['pplus'] = val
        pool = pools.get((lvl, role), {})
        pool_ok = len(pool.get('csw', [])) >= POOL_MIN
        # outing percentiles
        outing = {}
        for m, lab, _hi in OUTING_METRICS:
            v = cnt.get(m)
            outing[m] = {'value': v, 'pctl': pctl(pool.get(m, []), v) if (pool_ok and v is not None) else None}
        # per-type changes
        types_today = defaultdict(list); types_base = defaultdict(list)
        for r in today:
            types_today[r.get('Pitch Type')].append(r)
        for r in base:
            types_base[r.get('Pitch Type')].append(r)
        debut = len(base) == 0
        changes, new_pitches = [], []
        for pt, tp in types_today.items():
            if not pt:
                continue
            bp = types_base.get(pt, [])
            if not bp and not debut and len(tp) >= NEW_PITCH_MIN:
                m_v, _ = _mean_sd([x for x in (sf(p.get('Velocity')) for p in tp) if x is not None])
                new_pitches.append({'pt': pt, 'n': len(tp), 'velo': m_v})
                continue
            if debut:
                continue
            rec = type_deltas(tp, bp, pt, hand_sign, cnt['n'], len(base))
            changes.append(rec)
        flags = []
        for np_ in new_pitches:
            flags.append({'kind': 'new', 'pt': np_['pt'], 'n': np_['n'], 'velo': np_['velo'],
                          'text': (f"New pitch: {np_['n']} {PITCH_NAMES.get(np_['pt'], np_['pt'])}"
                                   f"{'s' if np_['n'] != 1 else ''} ({np_['pt']}), "
                                   f"{'%.1f mph' % np_['velo'] if np_['velo'] is not None else 'no velocity'}. "
                                   f"He had never thrown this tag before in the 2026 data.")})
        for key in [k for k, *_ in MEAN_COLS] + ['usage', 'whiff', 'chase']:
            hits = [(rec, rec['cols'][key]) for rec in changes if key in rec['cols']]
            pos = [h for h in hits if h[1]['z'] >= FLAG_Z]
            neg = [h for h in hits if h[1]['z'] <= -FLAG_Z]
            if key in ('whiff', 'chase'):
                chosen = [h for h in hits if abs(h[1]['z']) >= SINGLE_Z]
            elif key == 'usage':
                chosen = pos + neg if len(pos) + len(neg) >= 2 else [h for h in hits if abs(h[1]['z']) >= SINGLE_Z]
            else:
                chosen = pos if len(pos) >= 2 else (neg if len(neg) >= 2 else [h for h in hits if abs(h[1]['z']) >= SINGLE_Z])
            for rec, c in chosen:
                flags.append({'kind': 'change', 'pt': rec['pt'], 'col': key, 'z': c['z'],
                              'today': c['today'], 'base': c['base'], 'delta': c['delta'],
                              'low_confidence': key in LOW_CONFIDENCE_COLS,
                              'text': change_sentence(rec, key, c, hand)})
        hits = [m for m, _l, _h in OUTING_METRICS
                if outing[m]['pctl'] is not None and outing[m]['pctl'] >= OUTING_PCTL]
        outing_ok = (cnt['k'] >= K_HEADLINE or 'pplus' in hits or len(hits) >= OUTING_MIN_HITS)
        if cnt['k'] >= K_HEADLINE:
            flags.append({'kind': 'outing', 'metric': 'k', 'value': cnt['k'],
                          'text': f"{cnt['k']} strikeouts in {cnt['bf']} batters faced."})
        for m, lab, _hi in OUTING_METRICS:
            o = outing[m]
            if outing_ok and m in hits:
                flags.append({'kind': 'outing', 'metric': m, 'value': o['value'], 'pctl': o['pctl'],
                              'text': (f"{lab} {fmt_metric(m, o['value'])}, "
                                       f"{ordinal(o['pctl'])} percentile of 2026 {lvl} "
                                       f"{'start-length' if role == 'SP' else 'relief-length'} outings.")})
        if not flags:
            continue
        score = sum(abs(f.get('z', 0)) for f in flags) + sum(1 for f in flags if f['kind'] != 'change') * 2.5
        report['pitchers'].append({
            'pitcher': name, 'team': team, 'opp': opp, 'level': lvl, 'hand': hand, 'role': role,
            'line': cnt, 'outing': outing, 'pool_ok': pool_ok, 'debut': debut,
            'types': sorted(changes, key=lambda r: -r['n_today']),
            'new_pitches': new_pitches, 'flags': flags, 'score': score})
    report['pitchers'].sort(key=lambda p: -p['score'])
    return report


def ordinal(n):
    n = int(n)
    if 10 <= n % 100 <= 20:
        return f'{n}th'
    return f"{n}{ {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th') }"


def fmt_metric(m, v):
    if v is None:
        return '-'
    if m in ('whiff', 'chase', 'csw'):
        return f'{v * 100:.1f}%'
    if m in ('stuff', 'loc', 'pplus'):
        return f'{v:.0f}'
    if m == 'xrv100':
        return f'{v:.2f}'
    return f'{v}'


def fmt_val(c, v):
    if v is None:
        return '-'
    if c.get('pct'):
        return f'{v * 100:.0f}%'
    return f'{v:.{c["dec"]}f}'


def change_sentence(rec, key, c, hand):
    name = PITCH_NAMES.get(rec['pt'], rec['pt'])
    up = c['delta'] > 0
    d = abs(c['delta'])
    if c.get('pct'):
        amount = f'{d * 100:.0f} points'
        today, base = f"{c['today'] * 100:.0f}%", f"{c['base'] * 100:.0f}%"
    else:
        amount = f"{d:.{c['dec']}f} {c['unit']}"
        today, base = f"{c['today']:.{c['dec']}f}", f"{c['base']:.{c['dec']}f}"
    label = COL_LABEL[key]
    if key == 'rx':
        # hand-signed delta: positive = farther toward his arm side
        direction = 'toward his arm side' if up else 'toward his glove side'
        return (f"His {name} release X moved {amount} {direction} ({base} to {today} ft, z {c['z']:.1f})"
                f"{' - low confidence, extension-class persistence' if key in LOW_CONFIDENCE_COLS else ''}.")
    if key == 'hb':
        direction = 'toward his arm side' if up else 'toward his glove side'
        return f"His {name} horizontal break moved {amount} {direction} ({base} to {today} in, z {c['z']:.1f})."
    word = 'up' if up else 'down'
    tail = ' Low confidence: extension day effects rarely persist.' if key in LOW_CONFIDENCE_COLS else ''
    unit = '' if c.get('pct') else f" {c['unit']}"
    return f"His {name} {label} was {word} {amount} ({base} to {today}{unit}, z {c['z']:.1f}).{tail}"


# ── Output ───────────────────────────────────────────────────────────────
def write_docx(report, path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
    doc.add_heading(f"Daily pitcher scan: {report['scan_date']}", level=0)
    n_p = len(report['pitchers'])
    n_new = sum(1 for p in report['pitchers'] for f in p['flags'] if f['kind'] == 'new')
    n_chg = sum(1 for p in report['pitchers'] for f in p['flags'] if f['kind'] == 'change')
    n_out = sum(1 for p in report['pitchers'] for f in p['flags'] if f['kind'] == 'outing')
    doc.add_paragraph(
        f"{n_p} pitchers flagged: {n_new} new pitches, {n_chg} change flags, {n_out} outing flags. "
        f"A change flag needs two pitch types moved the same way at |z| >= {FLAG_Z:.0f} on his own game-to-game scale, "
        f"or one type at |z| >= {SINGLE_Z:.0f} (whiff and chase: |z| >= {SINGLE_Z:.0f} only). "
        f"Outing flags fire at the {ordinal(OUTING_PCTL)} percentile of 2026 outings at the same level and role, "
        f"or at {K_HEADLINE}+ strikeouts. Outing floor {MIN_OUTING} pitches; per-type baseline floor {BASELINE_MIN}. "
        f"Generated {report['generated']}.")
    thin = [k for k in ('MLB_SP', 'MLB_RP', 'ROC_SP', 'ROC_RP') if report['pool_sizes'].get(k, 0) < POOL_MIN]
    if thin:
        doc.add_paragraph(f"Outing pools under {POOL_MIN} outings, so no outing flags for them: {', '.join(thin)}.")
    rest = report['pitchers'][TOP_N:]
    for p in report['pitchers'][:TOP_N]:
        ln = p['line']
        doc.add_heading(f"{p['pitcher']} ({p['team']}{' vs ' + p['opp'] if p['opp'] else ''}, {p['hand'] or '?'}HP, {p['role']})", level=1)
        pp = ln.get('pplus')
        doc.add_paragraph(
            f"{ln['n']} pitches, {ln['bf']} BF, {ln['k']} K, {ln['bb']} BB, {ln['h']} H, {ln['hr']} HR. "
            f"Whiff {fmt_metric('whiff', ln['whiff'])}, chase {fmt_metric('chase', ln['chase'])}, CSW {fmt_metric('csw', ln['csw'])}, "
            f"Stuff+ {fmt_metric('stuff', ln['stuff'])}, Loc+ {fmt_metric('loc', ln['loc'])}, "
            f"xRV/100 {fmt_metric('xrv100', ln['xrv100'])}, Pitching+ {fmt_metric('pplus', pp)}."
            + (' Season debut in the data: no change flags possible.' if p['debut'] else ''))
        for f in p['flags']:
            doc.add_paragraph(f['text'], style='List Bullet')
        # outing percentile line
        if p['pool_ok']:
            parts = [f"{lab} {ordinal(p['outing'][m]['pctl'])}" for m, lab, _ in OUTING_METRICS
                     if p['outing'][m]['pctl'] is not None]
            doc.add_paragraph('Outing percentiles (2026, same level and role): ' + ', '.join(parts) + '.')
        # per-type table
        cols = ['Type', 'N', 'Usage', 'Velo', 'RelX', 'RelZ', 'Ext', 'IVB', 'HB', 'Spin', 'Whiff%', 'Chase%']
        keys = [None, None, 'usage', 'velo', 'rx', 'rz', 'ext', 'ivb', 'hb', 'spin', 'whiff', 'chase']
        rows_ = p['types'] + [{'pt': n['pt'], 'n_today': n['n'], 'n_base': 0, 'cols': {}, 'new': True} for n in p['new_pitches']]
        if rows_:
            t = doc.add_table(rows=1, cols=len(cols)); t.style = 'Light Grid Accent 1'
            for i, h in enumerate(cols):
                t.rows[0].cells[i].text = h
            for rec in rows_:
                cells = t.add_row().cells
                cells[0].text = rec['pt'] + (' (new)' if rec.get('new') else '')
                cells[1].text = f"{rec['n_today']}/{rec['n_base']}"
                for i, k in enumerate(keys):
                    if k is None:
                        continue
                    c = rec['cols'].get(k)
                    if c is None:
                        cells[i].text = '-'; continue
                    txt = f"{fmt_val(c, c['today'])} ({'+' if c['delta'] > 0 else ''}{fmt_val(c, c['delta']).rstrip('%')}{'' if not c.get('pct') else ''})"
                    if abs(c['z']) >= FLAG_Z:
                        run = cells[i].paragraphs[0].add_run(txt); run.bold = True
                    else:
                        cells[i].text = txt
            doc.add_paragraph('N = today/season. Each cell: today (delta vs season). Bold = |z| >= 2. RelX and HB are hand-signed (positive = arm side).').runs[0].font.size = Pt(8)
    if rest:
        doc.add_heading('Also flagged', level=1)
        for p in rest:
            kinds = [f['text'] for f in p['flags']]
            doc.add_paragraph(f"{p['pitcher']} ({p['team']}, {p['role']}, {p['line']['n']} pitches, {p['line']['k']} K): "
                              + ' '.join(kinds), style='List Bullet')
    doc.add_paragraph(
        'Persistence of a change flag into the next start, 2021-2025: release X 84%, spin 58%, velocity 54%, '
        'release height 53%, extension 42%, movement 36-38%, against base rates of 27-36%. '
        'Movement and extension flags are the weakest evidence of a lasting change.').runs[0].font.size = Pt(8)
    if report['skipped']:
        doc.add_paragraph(f"Under the {MIN_OUTING}-pitch floor: " + ', '.join(f"{s['pitcher']} ({s['n']})" for s in report['skipped']) + '.').runs[0].font.size = Pt(8)
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--date', default=SCAN_DATE)
    ap.add_argument('--cached', action='store_true', default=USE_CACHE)
    ap.add_argument('--force', action='store_true', default=FORCE)
    ap.add_argument('--out-dir', default=OUT_DIR)
    a = ap.parse_args()
    os.makedirs(STATE_DIR, exist_ok=True)
    today = date.today().isoformat()
    if not a.force and os.path.exists(MARKER) and open(MARKER).read().strip() == today:
        print(f'already ran today ({today}); pass --force to run again')
        return 0
    rows = load_pitches(a.cached)
    scan_date = a.date or max(r['Game Date'] for r in rows)
    if not any(r['Game Date'] == scan_date for r in rows):
        print(f'no pitches on {scan_date}; nothing to scan'); return 1
    print(f'scan date {scan_date}')
    report = scan(rows, scan_date)
    jpath = os.path.join(STATE_DIR, f'scan_{scan_date}.json')
    tmp = jpath + '.tmp'
    with open(tmp, 'w') as fh:
        json.dump(report, fh, indent=1, default=float)
    os.replace(tmp, jpath)
    dpath = os.path.join(a.out_dir, f'daily_pitcher_scan_{scan_date}.docx')
    write_docx(report, dpath)
    with open(MARKER, 'w') as fh:
        fh.write(today)
    print(f"{len(report['pitchers'])} pitchers flagged; wrote {dpath} and {jpath}")
    return 0


if __name__ == '__main__':
    sys.exit(main())
